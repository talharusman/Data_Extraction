"""
Step 3: Assemble extracted_data.jsonl into a workbook with the reference layout.

Sheets:
1. INFO
2. MASTER_DATASET
3. HEADER_DICTIONARY
"""
from __future__ import annotations

import json
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pipeline_config import COLUMNS, OUT_JSONL, OUT_XLSX


TITLE_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
ALT_FILL = PatternFill("solid", fgColor="FFF8FBFF")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFFFF")
TITLE_FONT = Font(color="FFFFFFFF", bold=True, size=14)
HEADER_FONT = Font(color="FFFFFFFF", bold=True, size=11)
BODY_FONT = Font(color="FF000000", size=11)
THIN = Side(style="thin", color="FFB7C9E2")
HEADER_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_rows():
    rows = []
    with open(OUT_JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            rec = json.loads(line)
            rows.append(rec)
    return rows


def normalize_value(value):
    if value is None or value == "N/A":
        return None
    return value


def build_info_rows(rows):
    total = len(rows)
    lead_counts = Counter((row.get("LEAD_MARKER") or "").strip().upper() for row in rows)
    bnk = lead_counts.get("BNK", 0)
    ibg = lead_counts.get("IBG", 0)
    return [
        ("Purpose", "Description of dataset"),
        ("How to use", "Usage instructions"),
        ("Dynamic record count", total),
        ("Dynamic BNK count", bnk),
        ("Dynamic IBG count", ibg),
        ("Notes", "Additional notes"),
    ]


def style_title_row(ws, row_idx, end_col):
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=end_col)
    cell = ws.cell(row_idx, 1)
    cell.fill = TITLE_FILL
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_idx].height = 18


def style_header_row(ws, row_idx, end_col):
    for col in range(1, end_col + 1):
        cell = ws.cell(row_idx, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def apply_body_style(cell, use_alt=False):
    cell.fill = ALT_FILL if use_alt else WHITE_FILL
    cell.font = BODY_FONT
    cell.border = BODY_BORDER
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def autosize(ws, widths):
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_info_sheet(wb, rows):
    ws = wb.create_sheet("INFO")
    ws.sheet_view.showGridLines = False
    style_title_row(ws, 1, 2)
    ws["A1"] = "NBO MASTER DATASET"

    info_rows = build_info_rows(rows)
    for r_idx, (field, value) in enumerate(info_rows, start=2):
        ws.cell(r_idx, 1, field)
        ws.cell(r_idx, 2, value)
        ws.row_dimensions[r_idx].height = 18
        ws.cell(r_idx, 1).font = Font(bold=True, size=11)
        ws.cell(r_idx, 2).font = BODY_FONT
        ws.cell(r_idx, 1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(r_idx, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    return ws


def build_master_sheet(wb, rows):
    ws = wb.create_sheet("MASTER_DATASET")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    for c_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, c_idx, col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28.8

    for r_idx, rec in enumerate(rows, start=2):
        for c_idx, col_name in enumerate(COLUMNS, start=1):
            value = normalize_value(rec.get(col_name))
            cell = ws.cell(r_idx, c_idx, value)
            apply_body_style(cell, use_alt=((r_idx - 2) % 2 == 1))

    widths = {
        1: 30, 2: 12, 3: 28, 4: 18, 5: 20, 6: 24, 7: 18, 8: 20,
        9: 20, 10: 18, 11: 10, 12: 10, 13: 12, 14: 12, 15: 18, 16: 18,
        17: 18, 18: 20, 19: 18, 20: 14, 21: 14, 22: 14, 23: 16, 24: 22,
        25: 14, 26: 14, 27: 14, 28: 16, 29: 18, 30: 18, 31: 16, 32: 14,
        33: 16, 34: 16, 35: 18, 36: 14, 37: 14, 38: 18, 39: 18, 40: 16,
        41: 14, 42: 18, 43: 24, 44: 28, 45: 18, 46: 18, 47: 16, 48: 22,
        49: 22, 50: 20, 51: 20, 52: 14, 53: 22, 54: 22, 55: 20, 56: 16,
    }
    autosize(ws, widths)
    return ws


def build_dictionary_sheet(wb):
    ws = wb.create_sheet("HEADER_DICTIONARY")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "DATA DICTIONARY"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 18

    headers = ["COLUMN_NAME", "DESCRIPTION", "TYPE_HINT", "EXAMPLE"]
    for c_idx, header in enumerate(headers, start=1):
        cell = ws.cell(2, c_idx, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    dictionary_rows = {
        "PRODUCT_NAME": ("Normalized product or service name.", "Text / categorical", "Alfalah Car Ijarah"),
        "LEAD_MARKER": ("Lead marker from user (IBG or BNK).", "Text / categorical", "IBG"),
        "SOURCE_FILE_PRODUCT": ("Batch/source group name provided by user.", "Text / categorical", "Isl Consumer"),
        "PLAN_TYPE": ("High-level family such as Loan, Deposit, Savings, Card, Investment, Insurance, Service or Loyalty.", "Text / categorical", "Loan"),
        "TARGET_GOAL": ("Primary customer need/use case.", "Text / categorical", "Housing"),
        "CUSTOMER_TYPE": ("Broad eligible customer class.", "Text / categorical", "SME / Corporate"),
        "EMPLOYMENT_TYPE": ("Employment eligibility where applicable.", "Text / categorical", "Salaried/SEP"),
        "CUSTOMER_SEGMENT": ("Behavioral/demographic segment when explicitly identified.", "Text / categorical", "NRP"),
        "TARGET_SEGMENT": ("More specific positioning segment where stated.", "Text / categorical", "Financial Inclusion"),
        "SEGMENT_TIER": ("Priority/HNW/premium segment tag.", "Text / categorical", "Premium (HNW)"),
        "MIN_AGE": ("Minimum eligible age.", "Numeric", "18"),
        "MAX_AGE": ("Maximum eligible age.", "Numeric", "65"),
        "GENDER": ("Gender eligibility/focus.", "Text / categorical", "Female"),
        "IS_BANK_OFFERED": ("1 for bank-offered/relationship products in this normalized dataset.", "Numeric", "1"),
        "ACCOUNT_TYPE": ("Current, Savings, Wallet, Digital, etc.", "Text / categorical", "Current"),
        "CARD_TYPE": ("Debit, Credit, Virtual Debit, Premium etc.", "Text / categorical", "Credit (Premium)"),
        "CHANNEL": ("Primary servicing/onboarding channel.", "Text / categorical", "Mobile App"),
        "ELIGIBILITY_TYPE": ("Short eligibility logic text.", "Text / categorical", "CNIC + Biometric"),
        "SERVICE_TYPE": ("Operational/service type for non-core products.", "Text / categorical", "Payments"),
        "REWARD_TYPE": ("Reward mechanism for loyalty rows.", "Text / categorical", "Points"),
        "CURRENCY": ("Currency or set of currencies.", "Text / categorical", "PKR + FCY"),
        "CURRENCY_TYPE": ("PKR/FCY style label if used.", "Text / categorical", "FCY"),
        "MIN_BALANCE": ("Minimum opening or operating balance.", "Numeric", "1000"),
        "AVG_BALANCE_REQUIREMENT": ("Average balance requirement.", "Numeric", "50000"),
        "MIN_INCOME": ("Minimum PKR income where stated.", "Numeric", "50000"),
        "MIN_INCOME_USD": ("Minimum USD income where stated.", "Numeric", "3000"),
        "MIN_INVESTMENT": ("Minimum investment/placement amount.", "Numeric", "100K"),
        "MIN_CONTRIBUTION": ("Minimum premium/contribution.", "Numeric", "250000"),
        "LOAN_AMOUNT_RANGE": ("Loan size or facility range.", "Text / categorical", "200K-3M"),
        "COVERAGE_AMOUNT": ("Coverage amount or insured amount.", "Text / categorical", "50K-150K coverage"),
        "FINANCING_TYPE": ("Conventional/Islamic financing structure or instrument type.", "Text / categorical", "Mudarabah"),
        "DEPOSIT_PROFIT_TYPE": ("Profit basis or mode.", "Text / categorical", "Tier-based"),
        "DEPOSIT_PROFIT_FREQUENCY": ("Monthly, semi-annual, maturity etc.", "Text / categorical", "Monthly"),
        "TENURE": ("Readable tenor text.", "Text / categorical", "1-5 years"),
        "TENURE_OPTIONS": ("Structured tenor menu text.", "Text / categorical", "1M -> 5Y"),
        "MIN_TERM_YEARS": ("Minimum term in years when directly available.", "Numeric", "10"),
        "MAX_TERM_YEARS": ("Maximum term in years when directly available.", "Numeric", "25"),
        "BUSINESS_TENURE": ("Required business age/operating history.", "Text / categorical", ">=3 years"),
        "COLLATERAL_TYPE": ("Security/collateral type.", "Text / categorical", "Property Mortgage"),
        "EQUITY_REQUIREMENT": ("Borrower equity or margin requirement.", "Text / categorical", "30%"),
        "DBR_LIMIT": ("Debt burden ratio limit.", "Text / categorical", "<=40%"),
        "TRANSACTION_LIMIT": ("Usage/balance/transaction cap.", "Text / categorical", "1M monthly"),
        "SPECIAL_CONDITIONS": ("Residual qualifiers or important caveats.", "Text / categorical", "RDA required"),
        "PRODUCT_DESCRIPTION": ("Short summary of the product or plan.", "Text / categorical", "Savings plan with flexible deposits"),
        "PROVIDER_NAME": ("Bank, insurer, or product provider name.", "Text / categorical", "Bank Alfalah"),
        "PRODUCT_VARIANT_TIER": ("Tier, variant, or package label.", "Text / categorical", "Premier"),
        "PRICING_RATE": ("Rate, markup, margin, or pricing summary.", "Text / categorical", "6.5% p.a."),
        "FEES_AND_CHARGES": ("Concise fees and charges summary.", "Text / categorical", "Issuance fee 500 PKR"),
        "KEY_BENEFITS": ("Key benefits or value points.", "Text / categorical", "Free withdrawals, digital access"),
        "OPTIONAL_RIDERS": ("Optional riders or add-ons.", "Text / categorical", "Accidental cover rider"),
        "FREE_LOOK_PERIOD_DAYS": ("Free-look period in days.", "Numeric", "14"),
        "REQUIRED_DOCUMENTS": ("Required application documents.", "Text / categorical", "CNIC, income proof"),
        "CLAIMS_SERVICE_CONTACT": ("Claims or service contact details.", "Text / categorical", "Call center 111-111-111"),
        "KEY_EXCLUSIONS": ("Main exclusions or limitations.", "Text / categorical", "Pre-existing conditions excluded"),
        "TAX_ZAKAT_TREATMENT": ("Tax or zakat treatment note.", "Text / categorical", "Zakat applicable"),
        "PREMIUM_PAYMENT_FREQUENCY": ("Frequency at which premiums/contributions may be paid, when explicitly stated (e.g., Annual, Semi-Annual, Quarterly, Monthly).", "Text / categorical", "Annual"),
    }

    for r_idx, col_name in enumerate(COLUMNS, start=3):
        desc, type_hint, example = dictionary_rows[col_name]
        row = [col_name, desc, type_hint, example]
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = BODY_BORDER
            cell.font = BODY_FONT if c_idx != 1 else Font(size=11, bold=False)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = ALT_FILL if (r_idx - 3) % 2 == 0 else WHITE_FILL
        if col_name == "SPECIAL_CONDITIONS":
            ws.row_dimensions[r_idx].height = 28.8

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:D{ws.max_row}"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 28
    return ws


def main():
    rows = load_rows()

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    build_info_sheet(wb, rows)
    build_master_sheet(wb, rows)
    build_dictionary_sheet(wb)

    wb.save(OUT_XLSX)
    print(f"Wrote {len(rows)} rows to {OUT_XLSX}")


if __name__ == "__main__":
    main()
