from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.models import DocumentSection, ParsedDocument
from parsers.base import BaseParser


class DocxParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        try:
            doc = Document(str(path))
        except PackageNotFoundError as exc:
            raise ValueError(f"Skipping invalid DOCX file: {path.name}") from exc
        sections: list[DocumentSection] = []
        current_heading: str | None = None

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.lower().startswith("heading"):
                current_heading = text
                sections.append(DocumentSection(text=text, heading=current_heading, section_type="heading"))
            else:
                sections.append(DocumentSection(text=text, heading=current_heading, section_type="paragraph"))

        for table_index, table in enumerate(doc.tables):
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                rows.append(" | ".join(cells))
            sections.append(
                DocumentSection(
                    text="\n".join(rows),
                    heading=current_heading,
                    section_type="table",
                    metadata={"table_index": table_index},
                )
            )

        return ParsedDocument(
            source_path=path,
            folder_name=path.parent.name,
            product_name=path.stem,
            sections=sections,
        )
