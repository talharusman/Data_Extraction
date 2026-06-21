from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import ProductExtraction

logger = logging.getLogger(__name__)

MASTER_COLUMNS = [
    "PRODUCT_NAME",
    "LEAD_CO_MNE",
    "SOURCE_FILE_PRODUCT",
    "PLAN_TYPE",
    "TARGET_GOAL",
    "CUSTOMER_TYPE",
    "EMPLOYMENT_TYPE",
    "CUSTOMER_SEGMENT",
    "TARGET_SEGMENT",
    "SEGMENT",
    "MIN_AGE",
    "MAX_AGE",
    "GENDER",
    "BANK_CUSTOMER",
    "ACCOUNT_TYPE",
    "CARD_TYPE",
    "CHANNEL",
    "ELIGIBILITY_TYPE",
    "SERVICE_TYPE",
    "REWARD_TYPE",
    "CURRENCY",
    "CURRENCY_TYPE",
    "MIN_BALANCE",
    "AVG_BALANCE_REQUIREMENT",
    "MIN_INCOME",
    "MIN_INCOME_USD",
    "MIN_INVESTMENT",
    "MIN_CONTRIBUTION",
    "LOAN_AMOUNT_RANGE",
    "COVERAGE_AMOUNT",
    "FINANCING_TYPE",
    "PROFIT_TYPE",
    "PROFIT_FREQUENCY",
    "TENURE",
    "TENURE_OPTIONS",
    "MIN_TERM_YEARS",
    "MAX_TERM_YEARS",
    "BUSINESS_TENURE",
    "COLLATERAL_TYPE",
    "EQUITY_REQUIREMENT",
    "DBR_LIMIT",
    "TRANSACTION_LIMIT",
    "SPECIAL_CONDITIONS",
]

BINARY_FIELDS = {"BANK_CUSTOMER"}
NUMERIC_FIELDS = {
    "MIN_AGE",
    "MAX_AGE",
    "MIN_BALANCE",
    "AVG_BALANCE_REQUIREMENT",
    "MIN_INCOME",
    "MIN_INCOME_USD",
    "MIN_INVESTMENT",
    "MIN_CONTRIBUTION",
    "COVERAGE_AMOUNT",
    "MIN_TERM_YEARS",
    "MAX_TERM_YEARS",
    "BUSINESS_TENURE",
    "EQUITY_REQUIREMENT",
}

TITLE_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
ROW_FILL_A = PatternFill("solid", fgColor="FFFFFFFF")
ROW_FILL_B = PatternFill("solid", fgColor="FFF8FBFF")
WHITE_FONT = Font(color="FFFFFFFF", bold=True, size=14)
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
FIELD_FONT = Font(bold=True)
TITLE_FONT = Font(color="FFFFFFFF", bold=True, size=14)
THIN_SIDE = Side(style="thin", color="FFD9E2F3")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CELL_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)


class ExcelExportAgent:
    def __init__(self, columns: list[str], products_path: str, review_path: str | None = None, review_threshold: float = 0.75) -> None:
        self.columns = columns
        self.products_path = Path(products_path)
        self.review_path = Path(review_path) if review_path else None
        self.review_threshold = review_threshold

    def export(self, products: list[ProductExtraction]) -> None:
        self.products_path.parent.mkdir(parents=True, exist_ok=True)
        normalized_new_rows = [self._row(product) for product in products]
        combined_rows = self._merge_rows(normalized_new_rows)
        workbook = self._build_workbook(combined_rows)
        self._save_workbook(workbook, self.products_path)
        if self.review_path:
            review_rows = [row for product, row in zip(products, normalized_new_rows, strict=True) if product.needs_review]
            review_workbook = self._build_workbook(review_rows or normalized_new_rows[:0])
            self._save_workbook(review_workbook, self.review_path)

    def _merge_rows(self, new_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        existing_rows = self._load_existing_rows()
        combined = existing_rows + new_rows
        deduped: dict[str, dict[str, Any]] = {}
        for row in combined:
            key = self._row_key(row)
            if key in deduped:
                deduped[key] = self._merge_row_values(deduped[key], row)
            else:
                deduped[key] = row
        return list(deduped.values())

    @staticmethod
    def _merge_row_values(existing: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
        merged = dict(existing)
        for column, value in incoming.items():
            if ExcelExportAgent._has_value(value):
                merged[column] = value
            elif column not in merged:
                merged[column] = value
        return merged

    @staticmethod
    def _has_value(value: Any) -> bool:
        if value is None:
            return False
        if isinstance(value, str) and not value.strip():
            return False
        return True

    def _load_existing_rows(self) -> list[dict[str, Any]]:
        if not self.products_path.exists():
            return []
        try:
            frame = pd.read_excel(self.products_path, sheet_name="MASTER_DATASET")
        except Exception as exc:
            logger.warning("Could not read existing workbook %s: %s", self.products_path, exc)
            return []
        rows: list[dict[str, Any]] = []
        for record in frame.to_dict(orient="records"):
            row = {column: self._normalize_value(column, record.get(column), record) for column in MASTER_COLUMNS}
            rows.append(row)
        return rows

    def _row_key(self, row: dict[str, Any]) -> str:
        parts = [
            self._key_part(row.get("LEAD_CO_MNE")),
            self._key_part(row.get("SOURCE_FILE_PRODUCT")),
            self._key_part(row.get("PRODUCT_NAME")),
        ]
        return "|".join(parts)

    @staticmethod
    def _key_part(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip().lower()
        return text

    def _build_workbook(self, rows: list[dict[str, Any]]) -> Workbook:
        wb = Workbook()
        default = wb.active
        wb.remove(default)

        info_ws = wb.create_sheet("INFO")
        master_ws = wb.create_sheet("MASTER_DATASET")
        dict_ws = wb.create_sheet("HEADER_DICTIONARY")

        self._write_info_sheet(info_ws, rows)
        self._write_master_sheet(master_ws, rows)
        self._write_dictionary_sheet(dict_ws)

        wb.active = wb.sheetnames.index("INFO")
        return wb

    def _write_info_sheet(self, ws, rows: list[dict[str, Any]]) -> None:
        ws.merge_cells("A1:D1")
        ws["A1"] = "NBO MASTER DATASET"
        ws["A1"].font = TITLE_FONT
        ws["A1"].fill = TITLE_FILL
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        summary_rows = [
            ("Purpose", "Consolidated eligibility dataset for model training built from all processed product batches."),
            ("How to use", "Use MASTER_DATASET as the base training table. Blank cells indicate attributes not explicitly identified in the source batch."),
            ("Dynamic record count", len(rows)),
            ("Dynamic BNK count", sum(1 for row in rows if self._lead_code(row) == "BNK")),
            ("Dynamic IBG count", sum(1 for row in rows if self._lead_code(row) == "IBG")),
            ("Notes", "BNK/IBG are inferred from the product family and source folder. Re-run export to upsert matching products."),
        ]

        for idx, (field, value) in enumerate(summary_rows, start=2):
            ws.cell(idx, 1, field)
            ws.cell(idx, 2, value)
            ws.cell(idx, 1).font = FIELD_FONT
            ws.cell(idx, 1).alignment = CELL_ALIGNMENT
            ws.cell(idx, 2).alignment = CELL_ALIGNMENT
            ws.cell(idx, 1).border = THIN_BORDER
            ws.cell(idx, 2).border = THIN_BORDER
            ws.cell(idx, 1).fill = ROW_FILL_A if idx % 2 == 0 else ROW_FILL_B
            ws.cell(idx, 2).fill = ROW_FILL_A if idx % 2 == 0 else ROW_FILL_B

        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 110
        ws.column_dimensions["C"].width = 4
        ws.column_dimensions["D"].width = 4
        ws.row_dimensions[1].height = 24
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A2:B7"

    def _write_master_sheet(self, ws, rows: list[dict[str, Any]]) -> None:
        for col_idx, column in enumerate(self.columns, start=1):
            cell = ws.cell(1, col_idx, column)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(self.columns))}1"

        for row_idx, row in enumerate(rows, start=2):
            fill = ROW_FILL_A if row_idx % 2 == 0 else ROW_FILL_B
            for col_idx, column in enumerate(self.columns, start=1):
                value = self._normalize_value(column, row.get(column), row)
                cell = ws.cell(row_idx, col_idx, value)
                cell.alignment = CELL_ALIGNMENT
                cell.fill = fill
                cell.border = THIN_BORDER

        self._auto_size_columns(ws, self.columns, min_width=12, max_width=42)

    def _write_dictionary_sheet(self, ws) -> None:
        ws.merge_cells("A1:D1")
        ws["A1"] = "DATA DICTIONARY"
        ws["A1"].font = TITLE_FONT
        ws["A1"].fill = TITLE_FILL
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        headers = ["COLUMN_NAME", "DESCRIPTION", "TYPE_HINT", "EXAMPLE"]
        for idx, header in enumerate(headers, start=1):
            cell = ws.cell(2, idx, header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

        for row_idx, column in enumerate(self.columns, start=3):
            description, type_hint, example = self._column_meta(column)
            values = [column, description, type_hint, example]
            fill = ROW_FILL_A if row_idx % 2 == 1 else ROW_FILL_B
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row_idx, col_idx, value)
                cell.alignment = CELL_ALIGNMENT
                cell.fill = fill
                cell.border = THIN_BORDER

        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:D{ws.max_row}"
        self._auto_size_columns(ws, headers, min_width=14, max_width=45)
        ws.column_dimensions["B"].width = 48
        ws.column_dimensions["D"].width = 30

    def _auto_size_columns(self, ws, headers: list[str], min_width: int = 12, max_width: int = 40) -> None:
        for col_idx, header in enumerate(headers, start=1):
            max_len = len(str(header))
            for row_idx in range(1, ws.max_row + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is None:
                    continue
                max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, min_width), max_width)

    def _save_workbook(self, workbook: Workbook, path: Path) -> None:
        try:
            workbook.save(path)
        except PermissionError:
            fallback = path.with_name(f"{path.stem}_local{path.suffix}")
            logger.warning("Could not write %s because it is open or locked; writing %s instead.", path, fallback)
            workbook.save(fallback)

    def _row(self, product: ProductExtraction) -> dict[str, Any]:
        row = {column: None for column in self.columns}
        source_path = Path(product.source_path)
        source_family = self._clean_text(source_path.parent.name or product.source_file)
        source_family = self._smart_title(source_family)
        lead_code = self._lead_code(product)

        for column in self.columns:
            field = product.fields.get(column)
            value = field.value if field else None
            row[column] = self._normalize_value(column, value, row, product=product)

        row["PRODUCT_NAME"] = row.get("PRODUCT_NAME") or self._smart_title(product.product_id)
        row["LEAD_CO_MNE"] = lead_code
        row["SOURCE_FILE_PRODUCT"] = source_family
        return row

    def _normalize_value(
        self,
        column: str,
        value: Any,
        row: dict[str, Any],
        product: ProductExtraction | None = None,
    ) -> Any:
        if column == "PRODUCT_NAME":
            return self._normalize_product_name(value, product)
        if column == "LEAD_CO_MNE":
            return self._lead_code(product, value)
        if column == "SOURCE_FILE_PRODUCT":
            return self._source_file_product(product, value)
        if column in BINARY_FIELDS:
            return self._normalize_binary(value)
        if column in NUMERIC_FIELDS:
            return self._normalize_numeric(value)
        if column == "GENDER":
            return self._normalize_gender(value)
        if column == "PLAN_TYPE":
            return self._normalize_plan_type(value, product)
        if column == "TARGET_GOAL":
            return self._normalize_target_goal(value, product)
        if column == "CUSTOMER_TYPE":
            return self._normalize_customer_type(value)
        if column == "EMPLOYMENT_TYPE":
            return self._normalize_employment_type(value)
        if column in {"CUSTOMER_SEGMENT", "TARGET_SEGMENT", "SEGMENT"}:
            return self._compact_label(value)
        if column == "ACCOUNT_TYPE":
            return self._normalize_account_type(value, product)
        if column == "CARD_TYPE":
            return self._normalize_card_type(value, product)
        if column == "CHANNEL":
            return self._normalize_channel(value)
        if column == "ELIGIBILITY_TYPE":
            return self._normalize_eligibility_type(value)
        if column == "SERVICE_TYPE":
            return self._normalize_service_type(value)
        if column == "REWARD_TYPE":
            return self._normalize_reward_type(value)
        if column == "CURRENCY":
            return self._normalize_currency(value)
        if column == "CURRENCY_TYPE":
            return self._normalize_currency_type(value)
        if column == "FINANCING_TYPE":
            return self._normalize_financing_type(value)
        if column == "PROFIT_TYPE":
            return self._normalize_profit_type(value)
        if column == "PROFIT_FREQUENCY":
            return self._normalize_profit_frequency(value)
        if column == "TENURE":
            return self._compact_tenure(value)
        if column == "TENURE_OPTIONS":
            return self._compact_label(value)
        if column == "COLLATERAL_TYPE":
            return self._normalize_collateral_type(value)
        if column == "SPECIAL_CONDITIONS":
            return self._normalize_special_conditions(value)
        return self._compact_label(value)

    def _normalize_product_name(self, value: Any, product: ProductExtraction | None) -> Any:
        text = self._compact_label(value)
        if text:
            return text
        if product:
            return self._smart_title(product.product_id)
        return None

    def _source_file_product(self, product: ProductExtraction | None, value: Any) -> Any:
        if product:
            source_path = Path(product.source_path)
            folder = self._smart_title(self._clean_text(source_path.parent.name))
            if folder and folder.lower() != "documents":
                return folder
        return self._compact_label(value)

    def _lead_code(self, product: ProductExtraction | None = None, value: Any = None) -> Any:
        candidate_parts: list[str] = []
        if value is not None:
            candidate_parts.append(self._clean_text(value))
        if product:
            if isinstance(product, dict):
                candidate_parts.append(self._clean_text(product.get("source_file")))
                candidate_parts.append(self._clean_text(product.get("source_path")))
                candidate_parts.append(self._clean_text(product.get("SOURCE_FILE_PRODUCT")))
                candidate_parts.append(self._clean_text(product.get("PRODUCT_NAME")))
            else:
                candidate_parts.append(self._clean_text(product.source_file))
                candidate_parts.append(self._clean_text(product.source_path))
        haystack = " ".join(candidate_parts).lower()
        if not haystack:
            return self._compact_label(value)
        ibg_markers = [" islamic ", "takaful", "ibg", "shariah", "banca", "roshan digital", "premier islamic", "rda islamic", "islamic banking"]
        if any(marker.strip() in haystack for marker in ibg_markers):
            return "IBG"
        return "BNK"

    def _normalize_binary(self, value: Any) -> int | None:
        if value is None:
            return None
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if value in (0, 1):
                return int(value)
        text = self._clean_text(value).lower()
        if not text:
            return None
        positive = [
            "existing customer",
            "existing customers",
            "bank customer",
            "bank customers",
            "current customer",
            "current customers",
            "account holder",
            "account holders",
            "yes",
            "required",
        ]
        negative = [
            "not required",
            "no",
            "non-customer",
            "non customer",
            "open to all",
            "all customers",
            "new customer",
            "new customers",
        ]
        if any(token in text for token in negative):
            return 0
        if "bank alfalah" in text and "customer" in text:
            return 1
        if any(token in text for token in positive):
            return 1
        return None

    def _normalize_numeric(self, value: Any) -> int | float | None:
        if value is None:
            return None
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            if value.is_integer():
                return int(value)
            return value
        text = self._clean_text(value)
        if not text:
            return None
        text = text.replace(",", "")
        number_match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
        if number_match:
            number = float(number_match.group())
            if number.is_integer():
                return int(number)
            return number
        return self._compact_label(text)

    def _normalize_gender(self, value: Any) -> Any:
        text = self._clean_text(value).lower()
        if not text:
            return None
        if "male" in text and "female" in text:
            return "Any"
        if "female" in text:
            return "F"
        if "male" in text:
            return "M"
        if "any" in text:
            return "Any"
        return self._compact_label(value)

    def _normalize_customer_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("SME", ["sme", "small and medium", "small-medium"]),
            ("Corporate", ["corporate", "business banking", "institutional"]),
            ("Business", ["business", "merchant", "sole proprietor", "entrepreneur"]),
            ("Individual", ["individual", "personal", "retail"]),
            ("Women", ["women", "ladies", "female"]),
            ("Senior Citizen", ["senior citizen", "retired", "pensioner"]),
            ("Student", ["student", "students"]),
            ("NRP", ["nrp", "non resident", "non-resident", "roshan digital"]),
            ("Any", ["anyone", "all customers", "open to all"]),
        ], fallback=value)

    def _normalize_employment_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Salaried", ["salaried", "salary"]),
            ("Self-Employed", ["self-employed", "self employed", "business owner", "entrepreneur"]),
            ("Business", ["business", "merchant"]),
            ("Retired", ["retired", "pensioner"]),
            ("Student", ["student"]),
            ("Any", ["any", "all"]),
        ], fallback=value)

    def _normalize_plan_type(self, value: Any, product: ProductExtraction | None = None) -> Any:
        return self._first_keyword(
            value or (product.product_id if product else None),
            [
                ("Savings", ["savings", "saving", "profit account", "amdani", "munafa"]),
                ("Current Account", ["current account", "current"]),
                ("Term Deposit", ["term deposit", "fixed deposit", "deposit"]),
                ("Card", ["credit card", "debit card", "card"]),
                ("Loan", ["loan"]),
                ("Financing", ["finance", "financing"]),
                ("Protection", ["protection", "takaful", "insurance"]),
                ("Digital Banking", ["digital", "wallet", "mobile app", "internet banking"]),
            ],
            fallback=value,
        )

    def _normalize_target_goal(self, value: Any, product: ProductExtraction | None = None) -> Any:
        return self._first_keyword(
            value or (product.product_id if product else None),
            [
                ("Education", ["education", "school", "student"]),
                ("Health", ["health", "hospital", "medical"]),
                ("Protection", ["protection", "insurance", "takaful"]),
                ("Remittance", ["remittance", "money transfer"]),
                ("Payroll", ["payroll", "salary"]),
                ("Investment", ["investment", "invest"]),
                ("Saving", ["saving", "savings", "profit"]),
                ("Daily Banking", ["daily banking", "everyday banking"]),
                ("Business", ["business", "merchant", "commerce"]),
                ("Housing", ["home", "house", "housing", "apna ghar"]),
                ("Auto", ["car", "auto", "vehicle"]),
                ("General", ["general", "all purpose", "everyday"]),
            ],
            fallback=value,
        )

    def _normalize_account_type(self, value: Any, product: ProductExtraction | None = None) -> Any:
        return self._first_keyword(
            value or (product.product_id if product else None),
            [
                ("Savings Account", ["savings account", "saving account"]),
                ("Current Account", ["current account"]),
                ("Term Deposit", ["term deposit", "deposit"]),
                ("Digital Account", ["digital account", "wallet"]),
                ("RDA", ["roshan digital", "rda"]),
                ("Business Account", ["business account", "merchant"]),
            ],
            fallback=value,
        )

    def _normalize_card_type(self, value: Any, product: ProductExtraction | None = None) -> Any:
        return self._first_keyword(
            value or (product.product_id if product else None),
            [
                ("Credit Card", ["credit card"]),
                ("Debit Card", ["debit card"]),
                ("Virtual Card", ["virtual card"]),
                ("Premier Card", ["premier card"]),
            ],
            fallback=value,
        )

    def _normalize_channel(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Branch", ["branch", "counter"]),
            ("Digital", ["digital", "online"]),
            ("Mobile App", ["mobile app", "app"]),
            ("Internet Banking", ["internet banking"]),
            ("ATM", ["atm"]),
            ("Call Center", ["call center", "phone banking", "helpline"]),
            ("Wallet", ["wallet"]),
        ], fallback=value)

    def _normalize_eligibility_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Existing Customer", ["existing customer", "bank customer", "account holder"]),
            ("New Customer", ["new customer", "non-customer"]),
            ("Open To All", ["open to all", "all customers", "anyone"]),
        ], fallback=value)

    def _normalize_service_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Savings", ["savings", "saving"]),
            ("Payments", ["payment", "payments"]),
            ("Transfers", ["transfer", "remittance"]),
            ("Protection", ["protection", "takaful", "insurance"]),
            ("Financing", ["finance", "financing", "loan"]),
            ("Digital", ["digital", "mobile", "internet"]),
            ("Card", ["card"]),
        ], fallback=value)

    def _normalize_reward_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Cashback", ["cashback"]),
            ("Points", ["points", "reward points"]),
            ("Waiver", ["waiver", "fee waiver"]),
            ("Profit", ["profit", "return"]),
            ("Protection", ["protection", "takaful", "insurance"]),
        ], fallback=value)

    def _normalize_currency(self, value: Any) -> Any:
        text = self._clean_text(value).upper()
        if not text:
            return None
        if "PKR" in text:
            return "PKR"
        if "USD" in text:
            return "USD"
        if "AED" in text:
            return "AED"
        if "FCY" in text or "FOREIGN" in text:
            return "FCY"
        return self._compact_label(text)

    def _normalize_currency_type(self, value: Any) -> Any:
        text = self._clean_text(value).lower()
        if not text:
            return None
        if "foreign" in text or "fcy" in text:
            return "FCY"
        if "local" in text or "pkr" in text:
            return "PKR"
        return self._compact_label(value)

    def _normalize_financing_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Murabaha", ["murabaha"]),
            ("Musharakah", ["musharakah"]),
            ("Ijarah", ["ijarah"]),
            ("Salam", ["salam"]),
            ("Istisna", ["istisna"]),
            ("Conventional", ["conventional"]),
            ("Islamic", ["islamic"]),
            ("Lease", ["lease"]),
        ], fallback=value)

    def _normalize_profit_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Fixed", ["fixed"]),
            ("Floating", ["floating"]),
            ("Variable", ["variable"]),
            ("Shariah Compliant", ["shariah", "islamic"]),
        ], fallback=value)

    def _normalize_profit_frequency(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("Monthly", ["monthly"]),
            ("Quarterly", ["quarterly"]),
            ("Half-Yearly", ["half-yearly", "semi annual", "semi-annual"]),
            ("Yearly", ["yearly", "annual", "annually"]),
            ("At Maturity", ["maturity"]),
        ], fallback=value)

    def _compact_tenure(self, value: Any) -> Any:
        text = self._clean_text(value)
        if not text:
            return None
        text = text.replace("years", "Y").replace("year", "Y").replace("months", "M").replace("month", "M")
        text = re.sub(r"\bfor\b", "", text, flags=re.I)
        return self._compact_label(text)

    def _normalize_collateral_type(self, value: Any) -> Any:
        return self._first_keyword(value, [
            ("None", ["no collateral", "without collateral", "unsecured", "none"]),
            ("Property", ["property", "mortgage", "house"]),
            ("Deposit", ["deposit", "lien"]),
            ("Pledge", ["pledge"]),
            ("Guarantee", ["guarantee", "surety"]),
            ("Salary", ["salary"]),
        ], fallback=value)

    def _normalize_special_conditions(self, value: Any) -> Any:
        text = self._clean_text(value)
        if not text:
            return None
        parts = [part.strip() for part in re.split(r"[;.|]| and | or ", text) if part.strip()]
        if not parts:
            parts = [text]
        compact = "; ".join(self._smart_title(part) for part in parts[:3])
        if len(compact) > 120:
            compact = " ".join(compact.split()[:16])
        return compact

    def _first_keyword(self, value: Any, choices: list[tuple[str, list[str]]], fallback: Any = None) -> Any:
        text = self._clean_text(value).lower()
        if not text:
            return self._compact_label(fallback)
        for label, keywords in choices:
            if any(keyword in text for keyword in keywords):
                return label
        return self._compact_label(fallback if fallback is not None else value)

    def _compact_label(self, value: Any) -> Any:
        text = self._clean_text(value)
        if not text:
            return None
        text = re.sub(r"\s+", " ", text)
        text = text.split("\n", 1)[0].strip()
        text = re.split(r"[;.,]", text)[0].strip()
        text = re.sub(r"\b(customers?|account holders?|holders?|members?|applicants?|for|with|subject to)\b", "", text, flags=re.I)
        text = re.sub(r"\s+", " ", text).strip()
        if not text:
            return None
        if len(text) > 80:
            text = " ".join(text.split()[:8])
        return self._smart_title(text)

    @staticmethod
    def _clean_text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if float(value).is_integer():
                return str(int(value))
            return str(value)
        return str(value).replace("\r", " ").replace("\n", " ").strip()

    @staticmethod
    def _smart_title(text: str) -> str:
        if not text:
            return text
        words: list[str] = []
        for token in re.split(r"(\s+)", text):
            if not token.strip():
                words.append(token)
                continue
            if token.isupper() and len(token) > 1:
                words.append(token)
            else:
                words.append(token[:1].upper() + token[1:].lower())
        return "".join(words).strip()

    def _column_meta(self, column: str) -> tuple[str, str, Any]:
        meta = {
            "PRODUCT_NAME": ("Normalized product name", "Text", "Alfalah Car Ijarah"),
            "LEAD_CO_MNE": ("Lead company / market code", "Code", "BNK"),
            "SOURCE_FILE_PRODUCT": ("Source folder or product family", "Text", "Banca Takaful"),
            "PLAN_TYPE": ("Product plan type", "Text", "Savings Account"),
            "TARGET_GOAL": ("Primary customer goal", "Text", "Education"),
            "CUSTOMER_TYPE": ("Target customer type", "Text", "Individual"),
            "EMPLOYMENT_TYPE": ("Employment classification", "Text", "Salaried"),
            "CUSTOMER_SEGMENT": ("Customer segment", "Text", "Retail"),
            "TARGET_SEGMENT": ("Target segment", "Text", "Senior Citizen"),
            "SEGMENT": ("General segment label", "Text", "Consumer"),
            "MIN_AGE": ("Minimum eligible age", "Numeric", 18),
            "MAX_AGE": ("Maximum eligible age", "Numeric", 60),
            "GENDER": ("Eligible gender", "Text", "Any"),
            "BANK_CUSTOMER": ("Existing bank customer required", "Binary (0/1)", 1),
            "ACCOUNT_TYPE": ("Account type", "Text", "Current Account"),
            "CARD_TYPE": ("Card type", "Text", "Debit Card"),
            "CHANNEL": ("Delivery or service channel", "Text", "Mobile App"),
            "ELIGIBILITY_TYPE": ("Eligibility category", "Text", "Existing Customer"),
            "SERVICE_TYPE": ("Service type", "Text", "Digital"),
            "REWARD_TYPE": ("Reward or benefit type", "Text", "Cashback"),
            "CURRENCY": ("Currency code or family", "Text", "PKR"),
            "CURRENCY_TYPE": ("Local or foreign currency", "Text", "PKR"),
            "MIN_BALANCE": ("Minimum balance requirement", "Numeric", 1000),
            "AVG_BALANCE_REQUIREMENT": ("Average balance requirement", "Numeric", 1000),
            "MIN_INCOME": ("Minimum income requirement", "Numeric", 50000),
            "MIN_INCOME_USD": ("Minimum income in USD", "Numeric", 180),
            "MIN_INVESTMENT": ("Minimum investment amount", "Numeric", 10000),
            "MIN_CONTRIBUTION": ("Minimum contribution amount", "Numeric", 1000),
            "LOAN_AMOUNT_RANGE": ("Loan amount range", "Text", "50,000-500,000"),
            "COVERAGE_AMOUNT": ("Coverage amount", "Numeric", 100000),
            "FINANCING_TYPE": ("Financing method", "Text", "Murabaha"),
            "PROFIT_TYPE": ("Profit or pricing type", "Text", "Fixed"),
            "PROFIT_FREQUENCY": ("Profit application frequency", "Text", "Monthly"),
            "TENURE": ("Tenure", "Text", "12 Months"),
            "TENURE_OPTIONS": ("Tenure options", "Text", "1, 2, 3 Years"),
            "MIN_TERM_YEARS": ("Minimum term in years", "Numeric", 1),
            "MAX_TERM_YEARS": ("Maximum term in years", "Numeric", 5),
            "BUSINESS_TENURE": ("Minimum business tenure", "Numeric", 2),
            "COLLATERAL_TYPE": ("Collateral type", "Text", "Property"),
            "EQUITY_REQUIREMENT": ("Equity requirement", "Numeric", 20),
            "DBR_LIMIT": ("Debt burden ratio limit", "Text", "50%"),
            "TRANSACTION_LIMIT": ("Transaction limit", "Text", "500,000"),
            "SPECIAL_CONDITIONS": ("Additional eligibility conditions", "Text", "Unique CNIC required"),
        }
        description, type_hint, example = meta.get(column, (column.replace("_", " ").title(), "Text", "Example"))
        return description, type_hint, example
